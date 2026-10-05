import io
import re
import datetime as dt
from abc import ABCMeta, abstractmethod
from typing import Iterable, Optional, Union

from bs4 import BeautifulSoup
from bs4.element import NavigableString, Tag
from celery.utils.log import get_task_logger
from dateutil.parser import parse as parse_date
from pydantic import parse_obj_as, ValidationError

from reservoirs.models import Reservoir
from services.schemes import (Gismeteo, Roshydromet, RP5,
                              RushydroSituation, Situation)
from services.utils import parser_info

logger = get_task_logger(__name__)


class AbstractParser(metaclass=ABCMeta):
    @classmethod
    @abstractmethod
    def parse(cls, *args, **kwargs):
        pass


class RushydroParser(AbstractParser):
    params = {
        'date': 'water-date',
        'level': 'water-level',
        'free_capacity': 'water-polemk',
        'inflow': 'water-pritok',
        'outflow': 'water-rashod',
        'spillway': 'water-sbros',
    }

    @classmethod
    def get_values(cls, data: Union[Tag, NavigableString]) -> Iterable:
        return zip(*(data[i].split(',')[::-1] for i in cls.params.values()))

    @classmethod
    def preprocessing(cls, values: Iterable) -> list[dict]:
        return [dict(zip(cls.params.keys(), i)) for i in values]

    @classmethod
    def parse(cls, page: str, reservoir: Reservoir) -> list[RushydroSituation]:
        soup = BeautifulSoup(page, 'html.parser')
        situations = []

        try:
            options = soup.find('div', {'data-river': 'Все реки'})
            option = options.find('option', string=reservoir.station_name)

            logger.info(f'{cls.__name__} parsed {reservoir.name}')

            if option:
                data = cls.preprocessing(cls.get_values(option))
                situations = parse_obj_as(list[RushydroSituation], data)

        except (ValueError, AttributeError, ValidationError, KeyError) as e:
            logger.error(f'{cls.__name__} {repr(e)}')

        finally:
            return situations


class EbvuDocxParser(AbstractParser):
    reservoir_names: dict = {
        'Саяно-Шушенское': 'sayano',
        'Майнское': 'mainsk',
        'Красноярское': 'kras',
        'Иркутское': 'irkutsk',
        'Братское': 'bratsk',
        'Усть-Илимское': 'ust-ilim',
        'Богучанское': 'boguch',
        'Усть-Хантайское': 'ust-hantay',
        'Курейское': 'kurey',
    }

    @staticmethod
    def parse_level(value: str) -> Optional[float]:
        try:
            return float(value.replace(',', '.').strip())
        except (ValueError, AttributeError):
            return None

    @staticmethod
    def parse_int(value: str) -> Optional[int]:
        try:
            return int(round(float(value.replace(',', '.').strip())))
        except (ValueError, AttributeError):
            return None

    @classmethod
    def parse_inflow(cls, value: str) -> Optional[int]:
        if not value or not value.strip():
            return None

        total = value.strip().split('/')[-1]
        return cls.parse_int(total)

    @classmethod
    def parse(cls, content: bytes, date: dt.date) -> dict[str, Situation]:
        from docx import Document

        try:
            doc = Document(io.BytesIO(content))
        except Exception as error:
            logger.error(f'{cls.__name__} {repr(error)}')
            return {}

        if not doc.tables:
            logger.error(f'{cls.__name__} no tables')
            return {}

        situations = {}

        for row in doc.tables[0].rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]

            if len(cells) < 6:
                continue

            name = cells[0].split('(')[0].strip()
            slug = cls.reservoir_names.get(name)

            if slug is None:
                continue

            level = cls.parse_level(cells[1])

            if level is None:
                logger.warning(f'{cls.__name__} no level for {name}')
                continue

            try:
                situations[slug] = Situation(
                    date=date,
                    level=level,
                    free_capacity=None,
                    inflow=cls.parse_inflow(cells[5]),
                    outflow=cls.parse_int(cells[3]),
                    spillway=cls.parse_int(cells[4]),
                )
            except ValidationError as error:
                logger.error(f'{cls.__name__} {repr(error)}')

        logger.info(f'{cls.__name__} parsed date {date}: {len(situations)}')

        return situations


class RP5Parser(AbstractParser):
    @staticmethod
    def get_headlines(first_row: Union[Tag, NavigableString]) -> list[str]:
        return [cell.text for cell in first_row.find_all('td')]

    @staticmethod
    def get_values(row: Tag) -> list[Optional[str]]:
        values = []

        for cell in row.find_all('td'):
            if cell.find('div'):
                values.append(cell.find('div').text)
            else:
                values.append(cell.text)

        return values

    @classmethod
    def preprocessing(cls, headlines: list, row: Tag) -> dict:
        return dict(zip(headlines[::-1], cls.get_values(row)[::-1]))

    @classmethod
    def get_observations(cls, table: Union[Tag, NavigableString]) -> list[dict]:  # noqa(E501)
        date_str = table.find('td', **{'class_': 'cl_dt'})
        date = parse_date(date_str.text, parserinfo=parser_info)

        logger.info(f'{cls.__name__} parsed date {date.date()}')

        headlines = cls.get_headlines(table.find('tr'))
        observations = []
        last_hour = 24

        for row in table.find('tbody').contents[1:]:
            hours = int(row.find('div', **{'class_': 'dfs'}).text)

            if hours > last_hour:
                break

            observation = cls.preprocessing(headlines, row)
            observation['date'] = date + dt.timedelta(hours=hours)
            observations.append(observation)

            last_hour = hours

        return observations

    @classmethod
    def parse(cls, page: str) -> list[RP5]:
        soup = BeautifulSoup(page, 'html.parser')
        archive_table = soup.find('table', id='archiveTable')

        if not archive_table:
            logger.error(f'{cls.__name__} no content')
            return []

        try:
            observations = cls.get_observations(archive_table)
            return parse_obj_as(list[RP5], observations)

        except (AttributeError, ValidationError, IndexError) as error:
            logger.error(f'{cls.__name__} {repr(error)}')
            return []


class GismeteoParser(AbstractParser):
    @classmethod
    def parse(cls, data: dict) -> list[Gismeteo]:
        try:
            return parse_obj_as(list[Gismeteo], data['response'])
        except (KeyError, ValidationError) as error:
            logger.error(f'{cls.__name__} {repr(error)}')
            return []


class RoshydrometParser(AbstractParser):
    @staticmethod
    def get_values(row: Tag) -> list:
        return re.findall(r'сегодня?|-?[0-9]+|штиль?', row.text, re.I)[-5:]

    @classmethod
    def get_observations(cls, table: Union[Tag, NavigableString]) -> list[dict]:  # noqa(E501)
        keys = (
            'temp',
            'pressure',
            'precipitation',
            'cloudiness',
            'wind_speed',
            'date',
        )
        hours = {
            'ночью': 1,
            'днем': 13,
        }

        date = dt.date.today()
        forecasts = []

        for row in table.find_all('tr'):
            date_str = row.find('div', **{'class': 'date'})

            if date_str and date_str.text != 'Сегодня':
                date = parse_date(date_str.text, parserinfo=parser_info)

            time = row.find('div', **{'class': 'small'}).text
            datetime = dt.datetime.combine(date, dt.time(hours.get(time, 0)))

            values = cls.get_values(row)
            values.append(datetime)

            forecasts.append(dict(zip(keys, values)))

        return forecasts

    @classmethod
    def parse(cls, page: str) -> list[Roshydromet]:
        soup = BeautifulSoup(page, 'html.parser')
        forecast_table = soup.find('tbody')

        if not forecast_table:
            logger.error(f'{cls.__name__} no content')
            return []

        try:
            forecasts = cls.get_observations(forecast_table)
            return parse_obj_as(list[Roshydromet], forecasts)

        except (AttributeError, ValidationError, IndexError) as error:
            logger.error(f'{cls.__name__} {repr(error)}')
            return []
